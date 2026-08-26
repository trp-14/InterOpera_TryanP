"""Fill report_template.xlsx with computed figures (BUILD_PLAN.md Step 6).

Never modifies the template in place — always opens a *fresh* copy of the
template and saves to a separate output path. Zero LLM imports (this module
lives under src/compute, subject to the same constraint as figures.py).
"""
from __future__ import annotations

from pathlib import Path

import openpyxl

# Row order in report_template.xlsx (row 1 = header) matches
# figures.FIGURE_ORDER exactly — each figure maps 1:1 to its template row.
FIGURE_ROW = {
    "sgs_allocation": 2,
    "mas_bills_allocation": 3,
    "ig_corporate_allocation": 4,
    "high_yield_allocation": 5,
    "fx_bonds_allocation": 6,
    "structured_credit_allocation": 7,
    "cash_allocation": 8,
    "aggregate_non_ig": 9,
    "single_issuer_concentration": 10,
    "gre_concentration": 11,
    "liquidity_ratio": 12,
    "portfolio_duration": 13,
    "portfolio_dv01": 14,
}

_COL_VALUE = 3
_COL_LIMIT = 4
_COL_UTILIZATION = 5
_COL_STATUS = 6
_COL_SOURCE = 7


def write_report(template_path: str | Path, report_dicts: list[dict], output_path: str | Path) -> None:
    workbook = openpyxl.load_workbook(template_path)
    worksheet = workbook.active

    for report in report_dicts:
        row = FIGURE_ROW.get(report["figure"])
        if row is None:
            raise ValueError(f"unknown figure {report['figure']!r} has no template row mapping")

        if report["status"] == "ERROR":
            worksheet.cell(row=row, column=_COL_VALUE, value="ERROR")
            worksheet.cell(row=row, column=_COL_LIMIT, value="n/a")
            worksheet.cell(row=row, column=_COL_UTILIZATION, value="n/a")
            worksheet.cell(row=row, column=_COL_STATUS, value="ERROR")
            worksheet.cell(row=row, column=_COL_SOURCE, value=report.get("error", ""))
            continue

        worksheet.cell(row=row, column=_COL_VALUE, value=report["value"])
        worksheet.cell(row=row, column=_COL_LIMIT, value=report["limit"])
        worksheet.cell(row=row, column=_COL_UTILIZATION, value=report["utilization"])
        worksheet.cell(row=row, column=_COL_STATUS, value=report["status"])

        citation = report["citation"]
        source = f"{report['graph_path']} → {citation['source_doc']} p.{citation['page']}" if citation else report["graph_path"]
        worksheet.cell(row=row, column=_COL_SOURCE, value=source)

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


if __name__ == "__main__":
    # Demo/manual-check entry point only. Takes the config path as an
    # argument rather than defaulting to one, so this file — like the rest
    # of src/compute — never names a firm (CLAUDE.md section 3.1).
    import sys

    from src.compute.figures import compute_all_figures
    from src.compute.formatters import render_all
    from src.config.loader import load_config
    from src.graph.builder import build_graph

    if len(sys.argv) != 2:
        raise SystemExit("usage: python -m src.compute.excel_report <path-to-config.yaml>")

    root = Path(__file__).resolve().parents[2]
    graph = build_graph(root / "sample_docs" / "sample_fund_guidelines.pdf", root / "sample_docs" / "sample_holdings.csv")
    config = load_config(sys.argv[1])
    reports = render_all(compute_all_figures(graph, config), config)

    output_path = root / "artifacts" / "demo" / "report.xlsx"
    write_report(root / "sample_docs" / "report_template.xlsx", reports, output_path)
    print(f"wrote {output_path.relative_to(root)}")
