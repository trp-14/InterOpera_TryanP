from pathlib import Path

import openpyxl
import pytest

from src.compute.excel_report import FIGURE_ROW, write_report
from src.compute.figures import compute_all_figures
from src.compute.formatters import render_all
from src.config.loader import load_config
from src.graph.builder import build_graph

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE_PATH = ROOT / "sample_docs" / "report_template.xlsx"


@pytest.fixture(scope="module")
def firm_a_reports() -> list[dict]:
    graph = build_graph(ROOT / "sample_docs" / "sample_fund_guidelines.pdf", ROOT / "sample_docs" / "sample_holdings.csv")
    config = load_config(ROOT / "config" / "firm_a.yaml")
    return render_all(compute_all_figures(graph, config), config)


def test_all_13_rows_populated(firm_a_reports, tmp_path):
    output_path = tmp_path / "report_firm_A.xlsx"
    write_report(TEMPLATE_PATH, firm_a_reports, output_path)

    wb = openpyxl.load_workbook(output_path)
    ws = wb.active

    assert len(FIGURE_ROW) == 13
    for figure_name, row in FIGURE_ROW.items():
        value = ws.cell(row=row, column=3).value
        limit = ws.cell(row=row, column=4).value
        utilization = ws.cell(row=row, column=5).value
        status = ws.cell(row=row, column=6).value
        source = ws.cell(row=row, column=7).value

        assert value is not None, f"row {row} ({figure_name}) Value is empty"
        assert limit is not None, f"row {row} ({figure_name}) Limit is empty"
        assert utilization is not None, f"row {row} ({figure_name}) Utilization is empty"
        assert status is not None, f"row {row} ({figure_name}) Status is empty"
        assert source, f"row {row} ({figure_name}) Source is empty"
        assert "→" in source and "sample_fund_guidelines.pdf" in source


def test_template_still_blank_after_write(firm_a_reports, tmp_path):
    output_path = tmp_path / "report_firm_A.xlsx"
    write_report(TEMPLATE_PATH, firm_a_reports, output_path)

    template_wb = openpyxl.load_workbook(TEMPLATE_PATH)
    template_ws = template_wb.active
    for row in FIGURE_ROW.values():
        for col in (3, 4, 5, 6, 7):
            assert template_ws.cell(row=row, column=col).value is None


def test_output_written_to_separate_file(firm_a_reports, tmp_path):
    output_path = tmp_path / "report_firm_A.xlsx"
    write_report(TEMPLATE_PATH, firm_a_reports, output_path)
    assert output_path.exists()
    assert output_path != TEMPLATE_PATH


def test_specific_cells_match_expected_values(firm_a_reports, tmp_path):
    output_path = tmp_path / "report_firm_A.xlsx"
    write_report(TEMPLATE_PATH, firm_a_reports, output_path)

    wb = openpyxl.load_workbook(output_path)
    ws = wb.active

    sgs_row = FIGURE_ROW["sgs_allocation"]
    assert ws.cell(row=sgs_row, column=3).value == "35.0%"
    assert ws.cell(row=sgs_row, column=6).value == "OK"

    cash_row = FIGURE_ROW["cash_allocation"]
    assert ws.cell(row=cash_row, column=6).value == "BREACH"

    single_issuer_row = FIGURE_ROW["single_issuer_concentration"]
    assert ws.cell(row=single_issuer_row, column=6).value == "AT LIMIT"
